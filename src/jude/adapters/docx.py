from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as _Doc
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

PARAGRAPH_SEP = "\n\n"


@dataclass
class DocxExtraction:
    """Flat text extracted from a DOCX, plus warnings about risky content.

    The `paragraphs` list preserves order; their join with PARAGRAPH_SEP gives
    the text that goes through the detection/redaction pipeline. The
    `core_properties` block is exposed separately so the user can decide
    whether to wipe it (author names often leak).
    """

    paragraphs: list[str] = field(default_factory=list)
    core_properties: dict[str, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return PARAGRAPH_SEP.join(self.paragraphs)


class DocxAdapter:
    """Read text from a .docx file and write back a redacted copy.

    Limitations of v0:
      * Run-level formatting *within* a paragraph is collapsed; the redacted
        output preserves paragraph boundaries but not bold/italic spans.
      * Comments and tracked changes are detected and surfaced as warnings,
        but not currently rewritten.
      * Headers and footers are extracted and rewritten.
    """

    @staticmethod
    def read(path: Path | str) -> DocxExtraction:
        doc: _Doc = Document(str(path))
        out = DocxExtraction()

        for p in _iter_paragraphs(doc):
            text = p.text
            if text.strip():
                out.paragraphs.append(text)

        cp = doc.core_properties
        out.core_properties = {
            "author": cp.author or "",
            "last_modified_by": cp.last_modified_by or "",
            "title": cp.title or "",
            "subject": cp.subject or "",
            "comments": cp.comments or "",
        }

        if _has_comments(doc):
            out.warnings.append(
                "Document contains review comments which are NOT redacted in v0. "
                "Inspect or remove them in Word before processing."
            )
        if _has_tracked_changes(doc):
            out.warnings.append(
                "Document contains tracked changes which are NOT redacted in v0. "
                "Accept or reject them in Word before processing."
            )
        if any(out.core_properties.values()):
            out.warnings.append(
                "Document core properties (author, last modified by, etc.) "
                "may contain identifying metadata. Consider clearing them."
            )

        return out

    @staticmethod
    def write_redacted(
        source_path: Path | str,
        target_path: Path | str,
        redacted_paragraphs: list[str],
        clear_metadata: bool = True,
    ) -> None:
        """Produce a redacted copy of the document.

        The body is rebuilt from `redacted_paragraphs` (which must align with
        the paragraphs returned by `read`). If `clear_metadata` is True, all
        core properties known to leak identity are blanked.
        """

        doc: _Doc = Document(str(source_path))
        live_paragraphs = [p for p in _iter_paragraphs(doc) if p.text.strip()]
        if len(live_paragraphs) != len(redacted_paragraphs):
            raise ValueError(
                f"Paragraph count mismatch: source has {len(live_paragraphs)}, "
                f"got {len(redacted_paragraphs)} replacements."
            )
        for p, new_text in zip(live_paragraphs, redacted_paragraphs, strict=True):
            _replace_paragraph_text(p, new_text)

        if clear_metadata:
            cp = doc.core_properties
            cp.author = ""
            cp.last_modified_by = ""
            cp.title = ""
            cp.subject = ""
            cp.comments = ""

        doc.save(str(target_path))


def _iter_paragraphs(doc: _Doc) -> Iterable[Paragraph]:
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_cell_paragraphs(cell)
    for section in doc.sections:
        if section.header is not None:
            yield from section.header.paragraphs
        if section.footer is not None:
            yield from section.footer.paragraphs


def _iter_cell_paragraphs(cell: _Cell) -> Iterable[Paragraph]:
    yield from cell.paragraphs
    for table in cell.tables:
        for row in table.rows:
            for inner_cell in row.cells:
                yield from _iter_cell_paragraphs(inner_cell)


def _replace_paragraph_text(p: Paragraph, new_text: str) -> None:
    """Replace a paragraph's text by clearing all runs and writing one new run.

    This loses run-level formatting. v0 tradeoff: simplicity > fidelity.
    """

    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    p.add_run(new_text)


def _has_comments(doc: _Doc) -> bool:
    body = doc.element.body
    for tag in ("commentRangeStart", "commentReference"):
        if body.findall(f".//{qn('w:' + tag)}"):
            return True
    return False


def _has_tracked_changes(doc: _Doc) -> bool:
    body = doc.element.body
    for tag in ("ins", "del", "moveFrom", "moveTo"):
        if body.findall(f".//{qn('w:' + tag)}"):
            return True
    return False
