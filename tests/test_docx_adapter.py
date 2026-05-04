from pathlib import Path

from docx import Document

from jude.adapters.docx import PARAGRAPH_SEP, DocxAdapter


def _make_doc(path: Path, paragraphs: list[str], author: str = "Jane Doe") -> None:
    doc = Document()
    doc.core_properties.author = author
    doc.core_properties.last_modified_by = author
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(path))


def test_read_extracts_paragraphs(tmp_path: Path):
    src = tmp_path / "in.docx"
    _make_doc(src, ["First paragraph.", "Second paragraph."])
    extraction = DocxAdapter.read(src)
    assert extraction.paragraphs == ["First paragraph.", "Second paragraph."]
    assert extraction.text == "First paragraph." + PARAGRAPH_SEP + "Second paragraph."


def test_read_reports_author_metadata_warning(tmp_path: Path):
    src = tmp_path / "in.docx"
    _make_doc(src, ["Hello"], author="Jane Doe")
    extraction = DocxAdapter.read(src)
    assert any("metadata" in w.lower() for w in extraction.warnings)
    assert extraction.core_properties["author"] == "Jane Doe"


def test_write_redacted_replaces_paragraphs_and_clears_metadata(tmp_path: Path):
    src = tmp_path / "in.docx"
    dst = tmp_path / "out.docx"
    _make_doc(src, ["Hello Amazon.", "Goodbye Microsoft."], author="Jane Doe")
    DocxAdapter.write_redacted(src, dst, ["Hello Org1.", "Goodbye Org2."])
    out = Document(str(dst))
    texts = [p.text for p in out.paragraphs if p.text.strip()]
    assert texts == ["Hello Org1.", "Goodbye Org2."]
    assert out.core_properties.author in ("", None)


def test_write_redacted_rejects_paragraph_count_mismatch(tmp_path: Path):
    src = tmp_path / "in.docx"
    dst = tmp_path / "out.docx"
    _make_doc(src, ["A", "B"])
    try:
        DocxAdapter.write_redacted(src, dst, ["only-one"])
    except ValueError as e:
        assert "Paragraph count mismatch" in str(e)
    else:
        raise AssertionError("expected ValueError")
