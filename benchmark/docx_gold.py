"""Word-highlight annotation protocol for the benchmark corpus.

The cheapest way for a lawyer to annotate a document is in Word: one
highlight colour per entity type, no tooling. This module turns such a
`.docx` into a gold corpus document, and does the reverse — writes
Jude's own detections into a copy as highlights — so the reviewer
corrects rather than annotates from scratch (an active-learning loop
with Word as the UI).

Legend (`HIGHLIGHT_TO_TYPE`):

    YELLOW        PERSON      BLUE      EMAIL
    BRIGHT_GREEN  ORG         RED       PHONE
    TURQUOISE     LOC         VIOLET    IBAN (and other account / registry numbers)
    PINK          CASE_REF    TEAL      URL
    DARK_RED      SECRET
    GRAY_25       PUBLIC — an institution, statute or court the reviewer
                  confirms must *not* be redacted. Feeds the
                  public-body over-redaction metric; never a gold span.

The extracted `text` is exactly what `jude.adapters.docx.DocxAdapter`
produces (same paragraph order, same filtering, same separator), so
gold offsets line up with what the pipeline sees.

    python -m benchmark.docx_gold prefill matter.docx matter.review.docx
    # … review in Word, fix highlights …
    python -m benchmark.docx_gold ingest matter.review.docx --id doc_021_share_purchase_en \
        --title "Share purchase agreement" --out benchmark/corpus
"""

from __future__ import annotations

import argparse
import json
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph

from jude.adapters.docx import PARAGRAPH_SEP, DocxAdapter, _iter_paragraphs

from .schema import GoldSpan

HIGHLIGHT_TO_TYPE: dict[str, str] = {
    "YELLOW": "PERSON",
    "BRIGHT_GREEN": "ORG",
    "TURQUOISE": "LOC",
    "PINK": "CASE_REF",
    "BLUE": "EMAIL",
    "RED": "PHONE",
    "VIOLET": "IBAN",
    "TEAL": "URL",
    "DARK_RED": "SECRET",
    "GRAY_25": "PUBLIC",
}
TYPE_TO_HIGHLIGHT: dict[str, str] = {t: c for c, t in HIGHLIGHT_TO_TYPE.items()}

PUBLIC = "PUBLIC"
_NO_COLOUR = {None, WD_COLOR_INDEX.AUTO, WD_COLOR_INDEX.INHERITED}
_EDGE_WS = " \t\n\r "


@dataclass
class IngestResult:
    text: str
    gold_spans: list[GoldSpan]
    public_spans: list[GoldSpan]
    language: str
    warnings: list[str] = field(default_factory=list)
    paragraph_count: int = 0


# --- reading ----------------------------------------------------------------


def _live_paragraphs(doc) -> list[Paragraph]:  # noqa: ANN001
    return [p for p in _iter_paragraphs(doc) if p.text.strip()]


def _colour_name(run) -> str | None:  # noqa: ANN001
    c = run.font.highlight_color
    return None if c in _NO_COLOUR else c.name


def _runs_with_colour(p: Paragraph) -> list[tuple[str, str | None]]:
    """(text, highlight colour name) per run, hyperlink runs included, in
    the order `Paragraph.text` concatenates them."""

    out: list[tuple[str, str | None]] = []
    for item in p.iter_inner_content():
        if isinstance(item, Hyperlink):
            out.extend((r.text, _colour_name(r)) for r in item.runs)
        else:
            out.append((item.text, _colour_name(item)))
    joined = "".join(t for t, _ in out)
    if joined != p.text:
        raise ValueError(
            "Run texts do not concatenate to the paragraph text; "
            "offsets would be wrong. Paragraph: " + p.text[:60] + "…"
        )
    return out


def _coloured_segments(runs: list[tuple[str, str | None]]) -> list[list]:
    """Contiguous same-colour stretches as [start, end, colour], offsets
    relative to the paragraph."""

    segs: list[list] = []
    pos = 0
    for text, colour in runs:
        if text and colour:
            if segs and segs[-1][2] == colour and segs[-1][1] == pos:
                segs[-1][1] = pos + len(text)
            else:
                segs.append([pos, pos + len(text), colour])
        pos += len(text)
    return segs


def _trim(text: str, start: int, end: int) -> tuple[int, int]:
    while start < end and text[start] in _EDGE_WS:
        start += 1
    while end > start and text[end - 1] in _EDGE_WS:
        end -= 1
    return start, end


def _merge_whitespace_separated(text: str, segs: list[list]) -> list[list]:
    """Merge same-colour segments separated only by whitespace — a
    reviewer who highlights word by word leaves the space bare."""

    out: list[list] = []
    for seg in segs:
        if out and out[-1][2] == seg[2]:
            gap = text[out[-1][1]:seg[0]]
            if gap and gap.strip() == "":
                out[-1][1] = seg[1]
                continue
        out.append(list(seg))
    return out


def ingest_docx(path: Path | str, language: str | None = None) -> IngestResult:
    doc = Document(str(path))
    paragraphs = _live_paragraphs(doc)

    texts: list[str] = []
    gold: list[GoldSpan] = []
    public: list[GoldSpan] = []
    warnings: list[str] = []
    offset = 0
    for p in paragraphs:
        p_text = p.text
        segs = _coloured_segments(_runs_with_colour(p))
        trimmed = []
        for s, e, colour in segs:
            s, e = _trim(p_text, s, e)
            if s < e:
                trimmed.append([s, e, colour])
        for s, e, colour in _merge_whitespace_separated(p_text, trimmed):
            surface = p_text[s:e]
            typ = HIGHLIGHT_TO_TYPE.get(colour)
            if typ is None:
                warnings.append(
                    f"Highlight colour {colour} has no meaning in the legend; "
                    f"ignored: {surface!r}"
                )
                continue
            span = GoldSpan(start=offset + s, end=offset + e, type=typ, text=surface)
            (public if typ == PUBLIC else gold).append(span)
        texts.append(p_text)
        offset += len(p_text) + len(PARAGRAPH_SEP)

    text = PARAGRAPH_SEP.join(texts)
    assert text == DocxAdapter.read(path).text, "extraction diverged from DocxAdapter"

    if language is None:
        from jude.detect.language import detect_language

        language = detect_language(text, supported=("en", "fr", "nl")) or "en"

    return IngestResult(
        text=text,
        gold_spans=gold,
        public_spans=public,
        language=language,
        warnings=warnings,
        paragraph_count=len(paragraphs),
    )


def to_corpus_dict(result: IngestResult, doc_id: str, title: str, notes: str = "") -> dict:
    def _spans(xs: Iterable[GoldSpan]) -> list[dict]:
        return [{"start": s.start, "end": s.end, "type": s.type, "text": s.text} for s in xs]

    return {
        "id": doc_id,
        "title": title,
        "language": result.language,
        "text": result.text,
        "gold_spans": _spans(result.gold_spans),
        "public_spans": _spans(result.public_spans),
        "notes": notes,
        "source": "docx-highlight",
    }


# --- writing ----------------------------------------------------------------


def _clear_paragraph_content(p: Paragraph) -> None:
    """Remove runs, hyperlinks and everything else but the paragraph
    properties. Hyperlinks survive as plain text; run formatting is
    lost — acceptable for a review copy."""

    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def _rebuild_runs(p: Paragraph, segments: list[tuple[str, str | None]]) -> None:
    _clear_paragraph_content(p)
    for text, colour in segments:
        if not text:
            continue
        run = p.add_run(text)
        if colour:
            run.font.highlight_color = WD_COLOR_INDEX[colour]


def prefill_docx(
    source: Path | str,
    target: Path | str,
    spans: Iterable,  # noqa: ANN001 — anything with .start/.end/.type
    public: Iterable = (),  # noqa: ANN001
) -> None:
    """Write `spans` (and `public` mentions, grey) into a copy of `source`
    as highlights. Offsets refer to `DocxAdapter.read(source).text`."""

    doc = Document(str(source))
    paragraphs = _live_paragraphs(doc)
    bounds: list[tuple[int, int]] = []
    pos = 0
    for p in paragraphs:
        bounds.append((pos, pos + len(p.text)))
        pos += len(p.text) + len(PARAGRAPH_SEP)

    wanted: list[tuple[int, int, str]] = []
    for s in spans:
        colour = TYPE_TO_HIGHLIGHT.get(s.type)
        if colour is None or s.type == PUBLIC:
            raise ValueError(f"No highlight colour for type {s.type!r}")
        wanted.append((s.start, s.end, colour))
    for s in public:
        wanted.append((s.start, s.end, TYPE_TO_HIGHLIGHT[PUBLIC]))

    per_para: dict[int, list[tuple[int, int, str]]] = {i: [] for i in range(len(paragraphs))}
    for start, end, colour in wanted:
        home = [i for i, (ps, pe) in enumerate(bounds) if ps <= start and end <= pe]
        if not home:
            raise ValueError(
                f"Span {start}-{end} straddles a paragraph boundary or lies outside the text"
            )
        per_para[home[0]].append((start - bounds[home[0]][0], end - bounds[home[0]][0], colour))

    for i, p in enumerate(paragraphs):
        local = sorted(per_para[i])
        if not local:
            continue
        text = p.text
        segments: list[tuple[str, str | None]] = []
        cur = 0
        for s, e, colour in local:
            if s < cur:  # overlaps the previous span — first one wins
                continue
            if s > cur:
                segments.append((text[cur:s], None))
            segments.append((text[s:e], colour))
            cur = e
        if cur < len(text):
            segments.append((text[cur:], None))
        _rebuild_runs(p, segments)

    doc.save(str(target))


# --- CLI --------------------------------------------------------------------


def _cmd_ingest(args: argparse.Namespace) -> None:
    result = ingest_docx(args.docx, language=args.language)
    title = args.title or Document(str(args.docx)).core_properties.title or Path(args.docx).stem
    d = to_corpus_dict(result, doc_id=args.id, title=title, notes=args.notes)
    out = Path(args.out) / f"{args.id}.json"
    if out.exists() and not args.force:
        raise SystemExit(f"{out} exists; pass --force to overwrite")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(d, indent=1, ensure_ascii=False), encoding="utf-8")
    by_type: dict[str, int] = {}
    for s in result.gold_spans:
        by_type[s.type] = by_type.get(s.type, 0) + 1
    print(f"{out}: {len(result.text)} chars, {result.paragraph_count} paragraphs, "
          f"language {result.language}, {len(result.gold_spans)} gold spans "
          f"{dict(sorted(by_type.items()))}, {len(result.public_spans)} public spans")
    for w in result.warnings:
        print(f"  warning: {w}")


def _cmd_prefill(args: argparse.Namespace) -> None:
    from .negative import public_mentions
    from .runners.jude_full import JudeFullRunner

    text = DocxAdapter.read(args.docx).text
    preds = JudeFullRunner().predict(text)
    public = [] if args.no_public else public_mentions(text, [])
    prefill_docx(args.docx, args.out, preds, public=public)
    print(f"{args.out}: {len(preds)} spans highlighted, {len(public)} public mentions in grey")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)

    ing = sub.add_parser("ingest", help="highlighted .docx → corpus JSON")
    ing.add_argument("docx", type=Path)
    ing.add_argument("--id", required=True, help="document id, e.g. doc_021_share_purchase_en")
    ing.add_argument("--title", default=None)
    ing.add_argument("--language", default=None, help="en / fr / nl (auto-detected if omitted)")
    ing.add_argument("--notes", default="")
    ing.add_argument("--out", type=Path, default=Path(__file__).parent / "corpus")
    ing.add_argument("--force", action="store_true")
    ing.set_defaults(func=_cmd_ingest)

    pre = sub.add_parser("prefill", help="write Jude's detections into a copy as highlights")
    pre.add_argument("docx", type=Path)
    pre.add_argument("out", type=Path)
    pre.add_argument("--no-public", action="store_true", help="do not grey out whitelisted public bodies")
    pre.set_defaults(func=_cmd_prefill)

    args = parser.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
