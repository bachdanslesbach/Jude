"""Reviewing the existing gold in Word.

`export` writes a corpus document as a .docx with its gold spans as
highlights (and whitelisted public mentions in grey); the reviewer
corrects it in Word and `ingest --force` writes it back, printing what
changed. Title and notes survive a re-ingest unless overridden.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from benchmark.schema import GoldDocument
from jude.adapters.docx import DocxAdapter


def _write_corpus_doc(path: Path, text: str, spans: list[tuple[str, str]], public=()) -> Path:  # noqa: ANN001
    def _sp(needle, typ):  # noqa: ANN001
        i = text.index(needle)
        return {"start": i, "end": i + len(needle), "type": typ, "text": needle}
    path.write_text(json.dumps({
        "id": path.stem, "title": "A title", "language": "en", "text": text,
        "gold_spans": [_sp(n, t) for n, t in spans],
        "public_spans": [_sp(n, "PUBLIC") for n in public],
        "notes": "SHOULD redact: Acme.",
    }, ensure_ascii=False), encoding="utf-8")
    return path


class TestExport:
    def test_export_writes_gold_as_highlights_and_public_in_grey(self, tmp_path: Path):
        from benchmark.docx_gold import export_docx, ingest_docx

        text = "The European Commission fined Acme SA.\n\nCounsel: Sophie Martin, Brussels."
        src = _write_corpus_doc(tmp_path / "doc_099_test_en.json", text,
                                [("Acme SA", "ORG"), ("Sophie Martin", "PERSON"), ("Brussels", "LOC")])
        out = tmp_path / "doc_099_test_en.docx"
        export_docx(src, out)

        assert DocxAdapter.read(out).text == text
        back = ingest_docx(out, language="en")
        assert [(s.type, s.text) for s in back.gold_spans] == [
            ("ORG", "Acme SA"), ("PERSON", "Sophie Martin"), ("LOC", "Brussels"),
        ]
        # The whitelist hit is greyed even though the JSON had no public_spans.
        assert [(s.type, s.text) for s in back.public_spans] == [("PUBLIC", "European Commission")]

    def test_export_keeps_annotator_public_spans(self, tmp_path: Path):
        from benchmark.docx_gold import export_docx, ingest_docx

        text = "Acme SA sued before the Tribunal de l'entreprise."
        src = _write_corpus_doc(tmp_path / "doc_098_test_en.json", text, [("Acme SA", "ORG")],
                                public=["Tribunal de l'entreprise"])
        out = tmp_path / "doc_098_test_en.docx"
        export_docx(src, out)
        back = ingest_docx(out, language="en")
        assert [s.text for s in back.public_spans] == ["Tribunal de l'entreprise"]

    def test_corpus_gold_spans_never_overlap(self):
        # Word cannot highlight two colours at one place, and the scorer
        # matches one-to-one: an overlapping gold pair is an annotation
        # error (doc_012 had 'RVK' inside 'RVK-2026-BPC-014').
        from benchmark.run import CORPUS_DIR

        for p in sorted(CORPUS_DIR.glob("doc_*.json")):
            gold = GoldDocument.from_json(p)
            ss = sorted(gold.gold_spans, key=lambda s: s.start)
            for a, b in zip(ss, ss[1:]):
                assert b.start >= a.end, (gold.id, a.text, b.text)

    def test_export_every_corpus_document_round_trips(self, tmp_path: Path):
        from benchmark.docx_gold import export_docx, ingest_docx
        from benchmark.run import CORPUS_DIR

        for p in sorted(CORPUS_DIR.glob("doc_*.json")):
            gold = GoldDocument.from_json(p)
            out = tmp_path / f"{gold.id}.docx"
            export_docx(p, out)
            back = ingest_docx(out, language=gold.language)
            assert back.text == gold.text, gold.id
            assert [(s.start, s.end, s.type) for s in back.gold_spans] == \
                [(s.start, s.end, s.type) for s in gold.gold_spans], gold.id


class TestReingest:
    def test_diff_reports_added_removed_and_retyped_spans(self, tmp_path: Path):
        from benchmark.docx_gold import diff_spans
        from benchmark.schema import GoldSpan

        text = "Acme SA and Zeta NV met Sophie Martin in Brussels."
        def g(n, t):  # noqa: ANN001
            i = text.index(n)
            return GoldSpan(i, i + len(n), t, n)
        before = [g("Acme SA", "ORG"), g("Zeta NV", "ORG"), g("Sophie Martin", "PERSON")]
        after = [g("Acme SA", "ORG"), g("Sophie Martin", "ORG"), g("Brussels", "LOC")]
        d = diff_spans(before, after)
        assert [s.text for s in d.added] == ["Brussels"]
        assert [s.text for s in d.removed] == ["Zeta NV"]
        assert [(s.text, o.type, s.type) for o, s in d.retyped] == [("Sophie Martin", "PERSON", "ORG")]

    def test_reingest_keeps_title_and_notes_and_reports_diff(self, tmp_path: Path, capsys):  # noqa: ANN001
        from benchmark.docx_gold import export_docx, main

        text = "The European Commission fined Acme SA.\n\nCounsel: Sophie Martin, Brussels."
        src = _write_corpus_doc(tmp_path / "doc_097_test_en.json", text,
                                [("Acme SA", "ORG"), ("Sophie Martin", "PERSON")])
        out = tmp_path / "doc_097_test_en.docx"
        export_docx(src, out)

        # Reviewer adds Brussels in Word.
        doc = Document(str(out))
        para = [p for p in doc.paragraphs if "Brussels" in p.text][0]
        for r in para.runs:
            if "Brussels" in r.text:
                before, _, after = r.text.partition("Brussels")
                r.text = before
                hl = para.add_run("Brussels")
                hl.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
                para.add_run(after)
        doc.save(str(out))

        main(["ingest", str(out), "--id", "doc_097_test_en", "--out", str(tmp_path), "--force"])
        printed = capsys.readouterr().out
        assert "+ LOC 'Brussels'" in printed

        updated = json.loads(src.read_text(encoding="utf-8"))
        assert updated["title"] == "A title"
        assert updated["notes"] == "SHOULD redact: Acme."
        assert [(s["type"], s["text"]) for s in updated["gold_spans"]] == [
            ("ORG", "Acme SA"), ("PERSON", "Sophie Martin"), ("LOC", "Brussels"),
        ]
        assert [s["text"] for s in updated["public_spans"]] == ["European Commission"]
